import streamlit as st
import pandas as pd
from pypdf import PdfReader
from pptx import Presentation
import dashscope
import os
import io

# ────────────────────────────────────────────────
#  1. 全局常量 & 缓存读取（程序启动时执行一次）
# ────────────────────────────────────────────────

@st.cache_data(show_spinner="正在加载内置知识库...")
def load_builtin_knowledge():
    # 假设两个文件都放在项目根目录
    try:
        # 原调研要点（如果还有）
        survey_df = pd.read_excel("调研要点.xlsx", sheet_name=0)
        survey_text = survey_df.to_string(index=False)

        # TBS-V2.xlsx （你的核心规则表）
        tbs = pd.read_excel("TBS-V2.xlsx", sheet_name="Sheet2")
        
        # 可以按需做结构化处理，例如按模块/类别分组
        tbs_grouped = tbs.groupby("模块")  # 或按 "章","节","禁止类 / 限制类 / 关注类" 分组
        
        tbs_by_category = {}
        for cat in ["禁止类", "限制类", "关注类", "部分有", "是", "否"]:
            subset = tbs[tbs["禁止类 / 限制类 / 关注类"].str.contains(cat, na=False)]
            tbs_by_category[cat] = subset.to_dict(orient="records")

        return {
            "survey_text": survey_text,
            "tbs_raw": tbs,
            "tbs_grouped": tbs_grouped,
            "tbs_by_category": tbs_by_category,
            "tbs_text": tbs.to_string(index=False)  # 如果想直接塞prompt也行
        }
    except Exception as e:
        st.error(f"加载内置知识库失败: {e}")
        st.stop()


# 程序启动时加载一次
BUILTIN_KNOWLEDGE = load_builtin_knowledge()

# ────────────────────────────────────────────────
#  2. 文件读取函数（用户上传的商业计划书）
# ────────────────────────────────────────────────

def extract_text_from_uploaded_file(uploaded_file):
    if uploaded_file is None:
        return ""
    file_ext = os.path.splitext(uploaded_file.name)[1].lower()
    bytes_data = uploaded_file.read()
    uploaded_file.seek(0)  # reset for future read

    if file_ext == ".pdf":
        reader = PdfReader(io.BytesIO(bytes_data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    elif file_ext == ".pptx":
        prs = Presentation(io.BytesIO(bytes_data))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)
    else:
        return ""

# ────────────────────────────────────────────────
#  3. 主界面
# ────────────────────────────────────────────────

st.set_page_config(page_title="科技项目合规与竞争力分析", layout="wide")
st.title("科技项目分析（基于 TBS-V2 规则库）")
st.caption("仅分析用户上传文档中明确提到的信息 + 内置 TBS-V2 规则，不引入外部未提及公司")

uploaded_file = st.file_uploader("上传项目资料（PDF / PPTX）", type=["pdf", "pptx"])

if uploaded_file:
    with st.spinner("正在提取文档内容..."):
        doc_text = extract_text_from_uploaded_file(uploaded_file)
    
    if not doc_text.strip():
        st.error("无法从文件中提取有效文本，请检查文件内容。")
    else:
        st.success("文档内容提取完成")
        # 可选：显示提取的前 800 字预览
        with st.expander("文档提取预览（前800字）"):
            st.text(doc_text[:800] + "...")

    if st.button("开始按模块分析（严格基于文档+TBS-V2）", type="primary"):
        with st.spinner("正在生成结构化分析报告..."):

            prompt = f"""你是一位非常严谨的项目评价助手，专注于科技企业投资/合规/竞争力分析。
请严格按照以下要求分析用户上传的文档。

## 限制（必须遵守）
- 只分析文档中**明确出现**的公司、产品、技术、市场、团队、财务、融资等信息
- 不得臆想、补充文档中未提及的团队成员、财务数据、融资规模、合规问题
- 竞争对比**仅限于**文档中明确提到的竞品/同行企业名称，不得自行引入其他公司
- 核心团队构成、财务指标、公司架构合规情况、融资规模及资金用途 → 只使用文档内容 + TBS-V2 中法律/合同/知识产权/公司治理相关规则进行对照评价
- 禁止输出任何融资建议、估值建议、建议融资金额

## 分析模块（必须严格按此顺序输出）
1. 产品技术
2. 市场
3. 行业竞争情况（需与文档中提到的竞品做对比）
4. 核心团队构成
5. 财务指标
6. 公司架构合规情况
7. 融资规模及资金用途

## 输出格式要求（每个模块独立）
### 模块名称
- 文档中关键事实总结（用 bullet points）
- 对照 TBS-V2 相关规则的评价（引用或概括 TBS 中的禁止/限制/关注类要求）
- 若涉及竞品对比，则明确列出文档中提到的竞品，并做客观对比
- 综合判断：【高/中/低】风险 / 竞争力 / 合规性

现在请开始分析。

用户上传文档内容：
{doc_text[:12000]}  （若太长可适当截断，但保留关键事实）

内置 TBS-V2 规则摘要（重点参考销售、市场、知识产权、合同、合规相关部分）：
{BUILTIN_KNOWLEDGE['tbs_text'][:8000]}  （实际可传入更多或结构化分组）
"""

            try:
                response = dashscope.Generation.call(
                    model="qwen-max",  # 或 qwen-plus / qwen-turbo 根据配额选择
                    prompt=prompt,
                    temperature=0.1,   # 降低创造性，提高严谨度
                    max_tokens=6000,
                    result_format="message"
                )
                analysis_result = response.output.choices[0].message.content
                st.markdown(analysis_result)
            except Exception as e:
                st.error(f"调用大模型失败：{e}")
